#!/usr/bin/env python3
"""Builds one Word file with all current legal documents, for legal review.

Reads the versioned documents in packages/shared/src/legal/docs (the same files the website shows)
and adds the appendices: comparison with Yad2's documents and open questions.
Usage: python3 packages/shared/scripts/legal_to_docx.py [output.docx]
"""
import json, pathlib, re, sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parents[1] / 'src' / 'legal'
OUT = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else 'docs/legal/WHATS-POPPIN-legal-review.docx')
COMPANY = json.loads((ROOT / 'company.json').read_text(encoding='utf-8'))
ORDER = ['terms', 'advertisers', 'privacy', 'accessibility']
FONT = 'Arial'


def latest(doc):
    files = sorted(ROOT.glob(f'docs/{doc}.v*.he.json'), key=lambda p: int(re.search(r'\.v(\d+)\.', p.name).group(1)))
    return json.loads(files[-1].read_text(encoding='utf-8'))


def company_line():
    parts = [COMPANY['ownerName'] or '[שם המפעיל: להשלים]',
             f"{COMPANY['businessType']} מס' {COMPANY['businessId'] or '[להשלים]'}",
             COMPANY['address'] or '[כתובת: להשלים]']
    return ', '.join(parts)


def fill(text):
    return (text.replace('{{company}}', company_line()).replace('{{owner}}', COMPANY['ownerName'] or '[שם רכז הנגישות: להשלים]')
            .replace('{{email}}', COMPANY['email']))


def rtl(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement('w:bidi'); b.set(qn('w:val'), '1'); pPr.append(b)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return p


def run(p, text, bold=False, size=11, color=None):
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = FONT
    rPr = r._r.get_or_add_rPr()
    fonts = rPr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rPr.append(fonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs'):
        fonts.set(qn(k), FONT)
    rtl_el = OxmlElement('w:rtl'); rtl_el.set(qn('w:val'), '1'); rPr.append(rtl_el)
    cs = OxmlElement('w:szCs'); cs.set(qn('w:val'), str(size * 2)); rPr.append(cs)
    if bold:
        bcs = OxmlElement('w:bCs'); rPr.append(bcs)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def para(d, text='', bold=False, size=11, style=None, color=None, space_after=6):
    p = d.add_paragraph(style=style) if style else d.add_paragraph()
    rtl(p)
    p.paragraph_format.space_after = Pt(space_after)
    rich(p, text, bold=bold, size=size, color=color)
    return p


def rich(p, text, bold=False, size=11, color=None):
    for part in re.split(r'(\*\*[^*]+\*\*)', fill(text)):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run(p, part[2:-2], bold=True, size=size, color=color)
        else:
            run(p, part, bold=bold, size=size, color=color)


def heading(d, text, size=16):
    p = para(d, text, bold=True, size=size, space_after=8)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.keep_with_next = True
    return p


def table(d, header, rows, widths=None):
    t = d.add_table(rows=1, cols=len(header))
    t.style = 'Table Grid'
    tblPr = t._tbl.tblPr
    bidi = OxmlElement('w:bidiVisual'); tblPr.append(bidi)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.paragraphs[0].text = ''
        rtl(c.paragraphs[0]); run(c.paragraphs[0], h, bold=True, size=10)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].paragraphs[0].text = ''
            rtl(cells[i].paragraphs[0]); rich(cells[i].paragraphs[0], v, size=10)
    d.add_paragraph()


def write_doc(d, data):
    heading(d, data['title'], size=20)
    y, m, dd = data['updated'].split('-')
    para(d, f"גרסה {data['version']} · עודכן ב-{int(dd)}.{int(m)}.{y}", size=10, color='55585F')
    heading(d, 'בקצרה', size=14)
    for line in data['summary']:
        para(d, line, style='List Bullet')
    heading(d, 'תוכן העניינים', size=14)
    for i, s in enumerate(data['sections'], 1):
        para(d, f'{i}. {s["title"]}', space_after=2)
    for i, s in enumerate(data['sections'], 1):
        heading(d, f'{i}. {s["title"]}', size=14)
        for j, item in enumerate(s['items'], 1):
            text = item if isinstance(item, str) else item['text']
            p = para(d, '')
            run(p, f'{i}.{j}  ', bold=True)
            rich(p, text)
            if not isinstance(item, str):
                for li in item['list']:
                    para(d, li, style='List Bullet')
    d.add_page_break()


YAD2 = [
    # topic, Yad2, before, now
    ('הגבלת גיל', 'מגיל 18 בלבד', 'לא היה', 'נוסף: הגלישה פתוחה לכולם, חשבון מפרסם מגיל 18 (שאלה פתוחה)'),
    ('הגדרות ומונחים', 'מוגדרים בפרקים שונים', 'לא היה', 'נוסף: סעיף "הגדרות" בראש כל מסמך (גולש, מפרסם, החברה, הפלטפורמה)'),
    ('פלטפורמה בלבד, אחריות המפרסם לתוכן', 'יש', 'משפט אחד', 'הורחב: פרק נפרד בתנאי השימוש ובהסכם המפרסמים, ומשפט קצר בממשק'),
    ('ויתור והגבלת אחריות (AS IS)', 'מפורט מאוד', 'משפט אחד', 'הורחב: תנאי שימוש סעיף 11, הסכם מפרסמים סעיף 13 עם תקרה'),
    ('שירותי צד שלישי וקישורים חיצוניים', 'יש', 'לא היה', 'נוסף: אפליקציות ניווט, אתר ואינסטגרם של המפרסם'),
    ('סדר הצגת תוצאות וחיפוש', 'יש (אלגוריתם, שדרוג מודעות)', 'לא היה', 'נוסף: "לוהט עכשיו", קרבה וקידום בתשלום'),
    ('מיקום ומיפוי', 'מאגר מפות של צד שלישי', 'לא היה', 'נוסף: מיקום ומרחק AS IS'),
    ('נתונים סטטיסטיים למפרסם', 'אינם מצג או התחייבות', 'לא היה', 'נוסף בהסכם המפרסמים'),
    ('פרסום בתשלום ומסלולים', 'מפורט', 'משפט על "ממומן"', 'נוסף: מה המפרסם מקבל, תשלום, סימון "ממומן"'),
    ('ביטול והחזרים', 'פרק נפרד + עמוד ביטול עסקה', 'לא היה', 'נוסף: לפי המדיניות שאושרה (24 שעות). עמוד ביטול נפרד: מומלץ בשלב ב׳'),
    ('זכות לסרב, לערוך ולהסיר פרסום', 'יש', 'משפט אחד', 'הורחב בהסכם המפרסמים'),
    ('תוכן אסור', 'רשימה ארוכה', 'משפט אחד', 'הורחב: "שימוש הוגן" ותנאים בהסכם המפרסמים'),
    ('איסור איסוף מידע אוטומטי (scraping)', 'יש', 'משפט אחד', 'נוסף במפורש'),
    ('הודעה והסרה (דיווח על הפרה)', 'יש, כולל שעות מענה', 'לא היה', 'נוסף: דיווח בעמוד "צרו קשר", מענה תוך 48 שעות'),
    ('קניין רוחני', 'הזכויות בפרסום עוברות לחברה', 'לא היה', 'נוסף: רישיון לא בלעדי במקום העברת זכויות (שאלה פתוחה)'),
    ('שיפוי ע"י המפרסם', 'יש', 'לא היה', 'נוסף בהסכם המפרסמים'),
    ('הצהרות המפרסם וחקיקה רלוונטית', 'רשימת חוקים', 'לא היה', 'נוסף: רישיון עסק, היתרים, הגנת הצרכן, מזון, נגישות'),
    ('זמינות ושינויים בשירות', 'יש', 'לא היה', 'נוסף בשני המסמכים'),
    ('דין וסמכות שיפוט', 'ישראל, תל אביב-יפו', 'היה', 'נשאר: ישראל, תל אביב-יפו'),
    ('מדיניות פרטיות: מידע טכני ונתוני שימוש', 'מפורט', 'חלקי', 'הורחב, בהפרדה בין גולשים למפרסמים'),
    ('פרטיות: עוגיות וכלי צד שלישי', 'רשימת כלים', 'לא היה', 'נוסף: אין עוגיות פרסום, רק אחסון חיוני'),
    ('פרטיות: הרשאות מכשיר (מיקום, התראות)', 'יש', 'חלקי', 'הורחב: מיקום רק בהסכמה, לא נשמר; התראות'),
    ('פרטיות: העברה לחו"ל', 'יש', 'לא היה', 'נוסף (שאלה פתוחה)'),
    ('פרטיות: קטינים', 'יש', 'לא היה', 'נוסף'),
    ('פרטיות: זכויות עיון, תיקון ומחיקה', 'יש', 'משפט אחד', 'הורחב: מענה תוך 30 יום, פנייה לרשות להגנת הפרטיות'),
    ('פרטיות: אבטחת מידע', 'יש', 'לא היה', 'נוסף'),
    ('פרטיות: משך שמירת מידע', 'יש', 'לא היה', 'נוסף (התקופות המדויקות: שאלה פתוחה)'),
    ('פרטיות: דיוור שיווקי', 'יש', 'לא היה', 'נוסף: אין דיוור לגולשים בלי הסכמה מפורשת'),
    ('נגישות: רכז נגישות וזמן מענה', 'שני אנשי קשר, 48 שעות', 'מייל בלבד, 14 יום', 'עודכן: רכז נגישות, 48 שעות, מידע בפורמט נגיש'),
    ('נגישות: פירוט תפריט הנגישות', 'יש', 'לא היה', 'נוסף'),
    ('שירות חיפוש AI, צ\'אט, מספרים וירטואליים', 'יש', '—', 'לא רלוונטי לפלטפורמה שלנו כרגע'),
    ('הסכם אבטחת מידע עם ספקים', 'מסמך נפרד', '—', 'לא נדרש מסמך ציבורי. מומלץ לבדוק את תנאי הספקים (Supabase, Vercel)'),
]

QUESTIONS = [
    'זהות המפעיל: הפלטפורמה מופעלת על ידי עוסק יחיד, וכרגע שמו, מספר העוסק והכתובת לא מוצגים באתר. האם חובה להציג אותם (למשל לפי חוק הגנת הצרכן לגבי מפרסמים שמשלמים, או בהצהרת הנגישות)? האם הסוג (פטור או מורשה) משנה משהו בניסוח, למשל בנושא מע"מ?',
    'אישור בצ\'קבוקס (Clickwrap): האם סימון תיבה, לצד תיעוד של פרטי העסק, תאריך ושעה, גרסת ההסכם וכתובת IP, ומשלוח עותק במייל, מספיקים כדי שההסכם יחייב את המפרסם?',
    'תקרת אחריות: האחריות מוגבלת לסכום ששולם עבור הקידום הספציפי, ובפרסום חינם אין אחריות כספית. האם זה עלול להיחשב "תנאי מקפח" לפי חוק החוזים האחידים, כשהמפרסמים הם עסקים קטנים?',
    'שיפוי: המפרסם משפה את החברה כולל שכר טרחת עורך דין, והחברה מתחייבת להודיע לו ולאפשר לו להשתתף בהגנה. האם ההיקף סביר?',
    'ביטול והחזרים: האם הוראות הביטול של חוק הגנת הצרכן (עסקת מכר מרחוק) חלות על מפרסם שהוא עוסק יחיד? האם מדיניות "24 שעות לפני" תקינה?',
    'קניין רוחני: בחרנו ברישיון לא בלעדי להצגת התוכן, ולא בהעברת זכויות כמו ביד2. האם הרישיון מספיק לארכיון ולשיתוף ברשתות חברתיות?',
    'גיל: הגלישה פתוחה לכולם, וחשבון מפרסם רק מגיל 18. האם צריך להגביל גם גולשים (למשל להתחברות במייל)?',
    'מיקום: המיקום נשלח לשרת רק לצורך החיפוש ולא נשמר. האם, לאור תיקון 13 לחוק הגנת הפרטיות, נתוני מיקום הם "מידע בעל רגישות מיוחדת" שדורש הסכמה מעבר לאישור בדפדפן?',
    'חובת יידוע: האם מדיניות הפרטיות מספיקה, או שצריך גם הודעה קצרה ליד הטפסים שבהם נאסף מידע (התחברות, טופס מפרסם)?',
    'משך שמירה: לכמה זמן לשמור את תיעוד האישור (כולל IP), פרטי תשלום ונתוני שימוש? 7 שנים?',
    'העברה לחו"ל: חלק מהמידע נשמר באיחוד האירופי ובארצות הברית. האם הניסוח הנוכחי מספיק?',
    'האם יש חובה לרשום מאגר מידע או למנות ממונה הגנת פרטיות, בהיקף פעילות כמו שלנו?',
    'שינוי מסמכים: לגולשים כתבנו "המשך הגלישה פירושו הסכמה", ומפרסמים מתבקשים לאשר מחדש. האם זה תקף?',
    'סימון "ממומן": האם התגית מספיקה לפי הדין בנושא פרסומת?',
    'נגישות: האם יש פטור או הקלה לעסק קטן בתקנות הנגישות לשירותי אינטרנט? האם צריך להציג שם וטלפון של רכז הנגישות?',
    'תרגום: המסמכים יתורגמו לערבית ולאנגלית אחרי האישור. האם מספיק לקבוע שהנוסח העברי מחייב?',
    'סמכות שיפוט בלעדית בתל אביב-יפו: האם מתאים, או שעדיף לפי מקום מושב העוסק?',
    'הצהרות המפרסם: האם רשימת הרישיונות וההיתרים (רישיון עסק, היתרי אירוע, בטיחות, מזון) מספיקה, או שחסר משהו שרלוונטי לפופ-אפים?',
]

STAGE_B = [
    'תיבת אישור (לא מסומנת מראש) בטופס הרשמת מפרסם, בהגשת פופ-אפ ובבקשת קידום. בלי סימון אי אפשר להמשיך.',
    'שמירת תיעוד האישור במסד הנתונים: פרטי בית העסק, תאריך ושעה, גרסה וכתובת IP.',
    'מייל אישור אוטומטי למפרסם, עם קישור לגרסה שאישר.',
    'בקשה לאשר מחדש כשההסכם מתעדכן, בכניסה הבאה לאזור המפרסמים.',
    'מומלץ: עמוד "ביטול עסקה" נפרד (בהשראת יד2), כשיתחיל תשלום באתר עצמו.',
]


def main():
    d = Document()
    section = d.sections[0]
    sectPr = section._sectPr
    bidi = OxmlElement('w:bidi'); sectPr.append(bidi)
    st = d.styles['Normal']; st.font.name = FONT; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:cs'), FONT)

    heading(d, 'WHATS POPPIN: המסמכים המשפטיים, לעיון', size=22)
    para(d, 'תודה רבה שהסכמת להציץ במסמכים ולעזור לי. זה עוזר לי מאוד.')
    para(d, 'כאן ריכזתי את הנוסחים המעודכנים (בעברית), בדיוק כפי שהם מופיעים באתר. בסוף יש שלושה נספחים קצרים: השוואה למסמכים של יד2, שבהם נעזרתי כהשראה; כמה שאלות שלא הייתי בטוח לגביהן; ומה עוד מתוכנן להיבנות באתר.')
    para(d, 'כל הערה תתקבל בשמחה, גם קטנה. אין צורך לעבור על הכל. אם יש זמן רק לחלק, הכי חשוב לי הסכם ההתקשרות עם מפרסמים והשאלות שבנספח ב.')
    para(d, f'המפעיל: {company_line()} · מייל: {COMPANY["email"]}', size=10, color='55585F')
    heading(d, 'מה במסמך', size=14)
    for i, name in enumerate(ORDER, 1):
        data = latest(name)
        para(d, f'{i}. {data["title"]} (גרסה {data["version"]})', space_after=2)
    para(d, 'נספח א: השוואה לתקנון ולמדיניות הפרטיות של יד2', space_after=2)
    para(d, 'נספח ב: שאלות שאשמח לשמוע את דעתך עליהן', space_after=2)
    para(d, 'נספח ג: מה עוד ייבנה באתר (שלב ב׳)', space_after=2)
    heading(d, 'המונחים', size=14)
    for line in ['"גולש": כל מי שנכנס לפלטפורמה וצופה בפופ-אפים. לא משלם.', '"מפרסם": בית עסק שמפרסם פופ-אפ, בתשלום או בחינם.',
                 '"החברה" / "אנחנו": המפעיל את WHATS POPPIN.', '"הפלטפורמה": האתר והאפליקציה יחד.']:
        para(d, line, style='List Bullet')
    para(d, 'משפט בממשק, בעמוד כל פופ-אפ ומעל רשימות התוצאות במפה ובחיפוש: "המידע נמסר על ידי בית העסק. האחריות על המוצרים והשירות היא של בית העסק בלבד."')
    d.add_page_break()

    for name in ORDER:
        write_doc(d, latest(name))

    heading(d, 'נספח א: השוואה למסמכי יד2', size=18)
    para(d, 'יד2 היא פלטפורמה שמחברת בין מפרסמים לגולשים, ולכן המבנה שלה רלוונטי. למדנו ממנה את המבנה ואת סוגי הסעיפים, וכתבנו הכל מחדש במילים שלנו. לא הועתק טקסט.')
    table(d, ['נושא', 'יד2', 'אצלנו לפני', 'אצלנו עכשיו'], YAD2)
    d.add_page_break()

    heading(d, 'נספח ב: שאלות שאשמח לשמוע את דעתך עליהן', size=18)
    para(d, 'אלה המקומות שבהם לא הייתי בטוח בניסוח. אם חלק מהשאלות דורשות בדיקה מעמיקה, נוכל לדבר עליהן בהמשך.')
    for i, q in enumerate(QUESTIONS, 1):
        p = para(d, '')
        run(p, f'{i}.  ', bold=True)
        rich(p, q)
    d.add_page_break()

    heading(d, 'נספח ג: מה עוד ייבנה באתר (שלב ב׳, אחרי רכישת דומיין)', size=18)
    para(d, 'הסכם המפרסמים כבר מתאר את המנגנונים הבאים. הם ייבנו באתר בשלב הבא, ועד אז בקשות קידום מטופלות במייל:')
    for line in STAGE_B:
        para(d, line, style='List Bullet')

    OUT.parent.mkdir(parents=True, exist_ok=True)
    d.save(OUT)
    print('wrote', OUT)


if __name__ == '__main__':
    main()
